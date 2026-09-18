"""
agent.py

Aca vive el "cerebro" del agente: arma el contexto del negocio,
llama a Gemini, interpreta la respuesta y decide si hay que
avisarle a un humano (escalar) o registrar un pedido.
"""

import json
import os
import random
import string
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import httpx
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
NEGOCIO_PATH = DATA_DIR / "negocio.json"
PRODUCTOS_PATH = DATA_DIR / "productos.json"
LOGS_PATH = DATA_DIR / "logs.json"
PEDIDOS_PATH = DATA_DIR / "pedidos.json"
PEDIDOS_XLSX_PATH = DATA_DIR / "pedidos.xlsx"

DEFAULT_GEMINI_MODEL = "gemini-3.1-flash-lite"

# Historial de conversacion por cliente (en memoria; alcanza para la demo).
# clave: id de sesion (telefono simulado o numero real de WhatsApp)
historiales: dict[str, list[dict]] = {}


class ItemPedido(BaseModel):
    producto: str
    cantidad: float


class Pedido(BaseModel):
    items: list[ItemPedido] = []
    total: float = 0


class ResultadoAgente(BaseModel):
    respuesta: str
    escalar: bool = False
    motivo: str = ""
    pedido: Optional[Pedido] = None


def leer_json(path: Path, valor_por_defecto):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return valor_por_defecto


def guardar_json(path: Path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def construir_system_prompt() -> str:
    negocio = leer_json(NEGOCIO_PATH, {})
    productos = leer_json(PRODUCTOS_PATH, [])

    catalogo = "\n".join(
        f"- {p['nombre']} ({p['categoria']}): ${p['precio']} | stock: {p['stock']} unidades"
        for p in productos
    )

    return f"""Sos el asistente virtual de "{negocio.get('nombre')}", un kiosco/almacen de barrio.
Atendes a clientes por WhatsApp. Hablá en español rioplatense, de forma breve, amable y natural,
como lo haría un empleado del kiosco. No uses lenguaje robótico ni te presentes como "modelo de IA".

DATOS DEL NEGOCIO:
- Dirección: {negocio.get('direccion')}
- Horario: {negocio.get('horario')}
- Formas de pago: {negocio.get('formasDePago')}
- Envíos: {negocio.get('envios')}

CATÁLOGO Y PRECIOS ACTUALES:
{catalogo}

REGLAS IMPORTANTES:
1. Solo respondé con información real del catálogo y los datos del negocio de arriba. Si te preguntan
   por un producto que NO está en el catálogo, decilo con honestidad y no inventes precio ni stock.
2. Si el cliente quiere hacer un pedido, confirmá los productos, cantidades y el total, y marcalo como pedido.
3. Marcá "escalar": true (para que un empleado humano intervenga) cuando:
   - El cliente pide algo que no podés resolver vos (reclamos, devoluciones, preguntas raras, precios
     especiales, algo fuera del catálogo).
   - El cliente parece enojado, urgente, o insiste en hablar con una persona.
   - No estás seguro de la respuesta.
4. Nunca confirmes definitivamente un pedido como "entregado" ni proceses pagos: solo tomá el pedido,
   avisá que un encargado lo va a confirmar, y marcá "pedido" con el detalle.

FORMATO DE RESPUESTA:
Respondé SIEMPRE y ÚNICAMENTE con un JSON válido, sin texto extra, con esta forma exacta:
{{
  "respuesta": "el mensaje de texto que le vas a mandar al cliente por WhatsApp",
  "escalar": true o false,
  "motivo": "si escalar es true, explicá brevemente por qué; si es false, dejalo vacío",
  "pedido": null o {{ "items": [{{"producto": "nombre", "cantidad": numero}}], "total": numero }}
}}"""


async def llamar_gemini(system_prompt: str, historial: list[dict]) -> ResultadoAgente:
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        return ResultadoAgente(
            respuesta=(
                "(El agente todavía no tiene configurada la GEMINI_API_KEY en el archivo .env. "
                "Agregala para que pueda responder.)"
            ),
            escalar=True,
            motivo="Falta configurar la API key de Gemini",
        )

    model = os.getenv("GEMINI_MODEL") or DEFAULT_GEMINI_MODEL
    url = (
        f"https://generativelanguage.googleapis.com/v1beta/models/"
        f"{model}:generateContent?key={api_key}"
    )

    body = {
        "systemInstruction": {"parts": [{"text": system_prompt}]},
        "contents": [
            {"role": m["role"], "parts": [{"text": m["texto"]}]} for m in historial
        ],
        "generationConfig": {
            "temperature": 0.4,
            "responseMimeType": "application/json",
        },
    }

    async with httpx.AsyncClient(timeout=45) as client:
        res = await client.post(url, json=body)

    if res.status_code != 200:
        raise RuntimeError(f"Gemini respondió {res.status_code}: {res.text}")

    data = res.json()
    try:
        texto = data["candidates"][0]["content"]["parts"][0]["text"]
    except (KeyError, IndexError):
        texto = ""

    try:
        parsed = json.loads(texto)
        return ResultadoAgente(
            respuesta=parsed.get("respuesta") or "Disculpá, no pude generar una respuesta.",
            escalar=bool(parsed.get("escalar")),
            motivo=parsed.get("motivo") or "",
            pedido=parsed.get("pedido"),
        )
    except (json.JSONDecodeError, TypeError):
        return ResultadoAgente(
            respuesta=texto or "Disculpá, tuve un problema para responder. Ya te contacta un encargado.",
            escalar=True,
            motivo="La respuesta del modelo no vino en formato válido",
        )


def registrar_log(entrada: dict):
    logs = leer_json(LOGS_PATH, [])
    logs.append(entrada)
    guardar_json(LOGS_PATH, logs)


def _abrir_o_crear_excel_pedidos():
    """Abre data/pedidos.xlsx si ya existe, o lo crea con encabezados."""
    if PEDIDOS_XLSX_PATH.exists():
        wb = load_workbook(PEDIDOS_XLSX_PATH)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Pedidos"
        ws.append(["ID", "Fecha", "Cliente", "Productos", "Total", "Estado"])
        anchos = {"A": 22, "B": 18, "C": 16, "D": 55, "E": 10, "F": 12}
        for columna, ancho in anchos.items():
            ws.column_dimensions[columna].width = ancho
    return wb, ws


def registrar_pedido_excel(pedido_id: str, cliente: str, pedido: Pedido, fecha_iso: str, estado: str = "pendiente"):
    """Agrega una fila nueva a data/pedidos.xlsx. Si falla, no rompe el
    flujo del agente (el JSON sigue siendo la fuente de verdad)."""
    try:
        wb, ws = _abrir_o_crear_excel_pedidos()
        productos_texto = "; ".join(f"{item.cantidad}x {item.producto}" for item in pedido.items)
        fecha_legible = datetime.fromisoformat(fecha_iso).strftime("%d/%m/%Y %H:%M")
        ws.append([pedido_id, fecha_legible, cliente, productos_texto, pedido.total, estado])
        wb.save(PEDIDOS_XLSX_PATH)
    except Exception as e:
        print(f"[agent] No se pudo actualizar pedidos.xlsx: {e}")


def registrar_pedido(cliente: str, pedido: Pedido):
    pedidos = leer_json(PEDIDOS_PATH, [])
    sufijo = "".join(random.choices(string.ascii_lowercase + string.digits, k=6))
    pedido_id = f"{int(time.time())}-{sufijo}"
    fecha_iso = datetime.now(timezone.utc).isoformat()
    pedidos.append(
        {
            "id": pedido_id,
            "cliente": cliente,
            "items": [item.model_dump() for item in pedido.items],
            "total": pedido.total,
            "estado": "pendiente",  # pendiente | confirmado | rechazado
            "fecha": fecha_iso,
        }
    )
    guardar_json(PEDIDOS_PATH, pedidos)
    registrar_pedido_excel(pedido_id, cliente, pedido, fecha_iso)


async def responder_mensaje(cliente: str, mensaje: str, canal: str) -> ResultadoAgente:
    """Punto de entrada del agente. Se usa tanto desde el WhatsApp simulado
    como desde el webhook real de Twilio."""

    historial = historiales.setdefault(cliente, [])
    historial.append({"role": "user", "texto": mensaje})

    system_prompt = construir_system_prompt()

    try:
        resultado = await llamar_gemini(system_prompt, historial)
    except Exception as err:  # noqa: BLE001 - queremos degradar con gracia en la demo
        print(f"[agent] Error llamando a Gemini: {err}")
        resultado = ResultadoAgente(
            respuesta="Uy, tuve un problema técnico. Ya le aviso a un encargado para que te responda.",
            escalar=True,
            motivo=f"Error llamando a Gemini: {err}",
        )

    historial.append({"role": "model", "texto": resultado.respuesta})
    # Evita que el historial crezca sin límite en una demo larga.
    if len(historial) > 20:
        del historial[: len(historial) - 20]

    if resultado.pedido:
        registrar_pedido(cliente, resultado.pedido)

    registrar_log(
        {
            "fecha": datetime.now(timezone.utc).isoformat(),
            "canal": canal,
            "cliente": cliente,
            "mensaje": mensaje,
            "respuesta": resultado.respuesta,
            "escalar": resultado.escalar,
            "motivo": resultado.motivo,
        }
    )

    return resultado
